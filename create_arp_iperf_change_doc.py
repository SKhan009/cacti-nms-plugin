from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/saimakhan/Documents/ChatGPT/Cacti NMS Project/ARP_iPerf_Before_After_Changes.docx"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def borders(table, color="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    border = tbl_pr.first_child_found_in("w:tblBorders")
    if border is None:
        border = OxmlElement("w:tblBorders")
        tbl_pr.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = border.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            border.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=9, font="Aptos"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def code_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.name = "Cascadia Mono"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(40, 40, 40)
    return paragraph


def comparison_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Inches(1.5), Inches(2.45), Inches(2.45)]
    for cell, label, width in zip(table.rows[0].cells, ["Area", "Before", "After"], widths):
        cell.width = width
        shade(cell, "263238")
        set_cell_text(cell, label, bold=True, color="FFFFFF", size=9)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, width, text in zip(cells, widths, row):
            cell.width = width
            shade(cell, "F4F6F7" if index % 2 else "FFFFFF")
            set_cell_text(cell, text, size=8.5)
    borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)
for style_name in ["Title", "Heading 1", "Heading 2"]:
    styles[style_name].font.name = "Aptos Display"
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(22)
styles["Heading 1"].font.size = Pt(15)
styles["Heading 2"].font.size = Pt(11)

title = doc.add_paragraph(style="Title")
title.add_run("ARP and iPerf3 Diagnostic Changes")
# Remove Word's built-in Title accent rule; hierarchy is carried by type and spacing.
title_props = title._p.get_or_add_pPr()
title_borders = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_borders.append(title_bottom)
title_props.append(title_borders)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
run = subtitle.add_run("Before and after reference for the NMS plugin")
run.italic = True
run.font.color.rgb = RGBColor(90, 90, 90)

doc.add_heading("Purpose", level=1)
doc.add_paragraph(
    "This document records the implementation changes made to the on-demand collector diagnostics. "
    "It shows the old behavior, the new behavior, the exact source files and line numbers, and the reason for each change. "
    "The changes affect the NMS plugin diagnostics only; they do not modify Cacti core polling, RRD files, SNMP credentials, or device configuration."
)

doc.add_heading("Change summary", level=1)
comparison_table(doc, [
    ("Collector ARP", "Ran ip neigh show to the selected device address. A loopback target such as 127.0.0.1 returned no rows even when the collector had learned other neighbours.", "Runs one direct ip neigh show command for the complete collector neighbour cache. It includes IPv4 and IPv6 rows and displays the command and output. There is no alternate command fallback."),
    ("iPerf3", "Started iperf3 immediately. When TCP 5201 had no listening server, iperf3 returned partial JSON followed by a confusing Bad file descriptor error.", "Checks TCP 5201 before starting iperf3. When no server is reachable, the page explains that an authorised iPerf3 server must be started on the endpoint."),
    ("Readiness text", "ARP was described only as reading the collector cache.", "The page states that ARP reads all cached IPv4 and IPv6 neighbours and that iPerf3 requires a server on TCP 5201."),
])

doc.add_heading("Edited files shown in the change", level=1)
doc.add_paragraph(
    "These are the three files changed for the collector ARP and iPerf3 update. The first two contain the implementation and page-level wording; the third contains the visible run-panel description."
)
comparison_table(doc, [
    ("plugins/nms/diagnostics.php", "Lines 73-81: readiness copy described the older ARP behavior and iPerf3 requirement.", "Lines 73-81: explains that ARP displays the full collector neighbour cache and that iPerf3 needs TCP 5201."),
    ("plugins/nms/includes/diagnostics.php", "Lines 233-249: ARP was filtered to the selected target; iPerf3 started without a TCP preflight.", "Lines 233-238 and 287-291: runs one direct full-cache command; lines 268-285: checks TCP 5201 and returns a clear unavailable result."),
    ("plugins/nms/templates/diagnostics/run.php", "Line 33: the run panel described ARP without clarifying that it reads the collector cache.", "Line 33: says Collector ARP lookup displays the full collector neighbour cache and describes the iPerf3 endpoint requirement."),
])

doc.add_heading("SNMPSim page cleanup", level=1)
doc.add_paragraph(
    "The separate FCAPS lab-scenarios panel was removed from the SNMP recordings page because it was not needed in the workflow. "
    "The SNMPSim status controls, upload flow, imported-record inventory, live-SNMP check, and Add device actions remain available."
)
doc.add_heading("Before", level=2)
code_paragraph(doc, "File: plugins/nms/templates/repository/import.php\nLines 105-143 (before cleanup)\n\n<section class=\"nms-panel nms-fcaps-lab\" id=\"fcaps-lab-scenarios\">\n  FCAPS lab scenarios panel and per-record scenario forms\n</section>")
doc.add_paragraph(
    "This rendered the FCAPS lab-scenarios section, including interface up/down, traffic pulse, UPS battery, and system-name controls. "
    "The section also displayed the statement that changes only the selected imported SNMPSim record."
)
doc.add_heading("After", level=2)
code_paragraph(doc, "File: plugins/nms/templates/repository/import.php\nLines 103-103 after cleanup\n\n<div class=\"nms-upload-dialog...\">\n\nThe FCAPS section is no longer rendered. The existing SNMPSim status and record-management sections continue below it.")
doc.add_paragraph(
    "This is a presentation-only removal. It does not change real devices, Cacti credentials, production polling, or the SNMPSim service controls that remain on the page."
)

doc.add_heading("SNMPSim FCAPS up/down on offline RHEL", level=1)
doc.add_paragraph(
    "The interface up/down value is written to the selected imported .snmprec record. "
    "The value is then served by SNMPSim on the next responder load. Offline RHEL installations commonly use the portable manual configuration; a responder already running in that mode keeps the old record in memory, which made the change appear not to work."
)
doc.add_heading("Before", level=2)
code_paragraph(doc, "File: plugins/nms/includes/snmprec.php\nLines 351-354 before this fix\n\nreturn $label . \". Reload is queued when managed SNMPSim is enabled; run or wait for the next Cacti poll to see the change.\";\n\nThe message did not distinguish managed systemd activation from manual activation, even though manual mode does not reload a running responder.")
doc.add_heading("After", level=2)
code_paragraph(doc, "File: plugins/nms/includes/snmprec.php\nLines 351-360 after this fix\n\n$settings = nms_snmpsim_config();\nif (($settings[\"activation\"] ?? \"\") === \"manual\") {\n    return $label . \". The record file was updated, but manual SNMPSim activation does not reload a running responder. Restart the responder on the configured collector, then run the next Cacti poll.\";\n}\nreturn $label . \". Reload is queued; the managed SNMPSim timer will restart the responder before the next Cacti poll.\";")
doc.add_paragraph(
    "The record edit remains limited to the selected simulator import. It never changes a real Cacti device or credentials. For automatic offline RHEL behavior, install the generated systemd bundle and enable both snmpsim.service and snmpsim-reload.timer; the timer consumes the queued reload marker. With manual activation, restart the administrator-owned responder after the change, then poll again."
)
code_paragraph(doc, "Files: plugins/nms/config.example.php and NMS-Generic-Offline-1.9.34/nms/snmpsim/README.md\n\nThe manual-mode warning now explains that a responder restart is required. The offline README documents the systemd timer path for automatic reloads and the same manual-mode limitation.")
doc.add_paragraph("For the managed offline RHEL setup, run these commands after installing the generated service files:")
code_paragraph(doc, "sudo systemctl enable --now snmpsim.service snmpsim-reload.timer\nsudo systemctl restart snmpsim.service")
doc.add_paragraph(
    "The first command starts the responder and the reload timer. The second command reloads the current record immediately. After that, an interface up/down change queues a marker and the timer restarts SNMPSim automatically."
)

doc.add_heading("Topology review: ARP and iPerf3", level=1)
doc.add_paragraph(
    "No topology link is created from ARP or iPerf3. They answer different questions: ARP shows which IP and MAC neighbours the collector has learned, and iPerf3 measures TCP throughput to an endpoint. Neither result identifies the remote switch, the remote Ethernet port, or a reliable device-to-port relationship."
)
doc.add_heading("Why these diagnostics are separate", level=2)
doc.add_paragraph(
    "No ARP or iPerf3 code is used to create topology links. Topology links come from LLDP/CDP neighbour evidence, matched with IF-MIB interface and port state. ARP and iPerf3 remain available as separate on-demand diagnostics."
)
doc.add_heading("Current dynamic topology flow", level=2)
code_paragraph(doc, "Files: plugins/nms/includes/topology/discovery.php and plugins/nms/includes/topology/canvas.php\n\nTopology discovery lines 22-56: load permitted Cacti devices and stored discovery snapshots, validate the device configuration hash, mark stale evidence, and reconcile current links.\n\nPort matrix lines 62-118: read IF-MIB interface admin/oper state and match only current resolved LLDP/CDP links.\n\nCanvas interface payload lines 79-115: expose the live interface state, physical-port evidence, and computed availability to the topology UI.")
doc.add_paragraph(
    "The topology therefore remains dynamic for live devices: LLDP/CDP identifies the remote device and port; IF-MIB identifies whether the local port is up, down, disabled, absent, or up without a neighbour. "
    "ARP and iPerf3 remain available from Protocol checks and do not create fabricated topology links."
)
doc.add_heading("Operational requirement", level=2)
doc.add_paragraph(
    "For a live topology, assign an LLDP or CDP discovery preset to each permitted Cacti device, allow the relevant neighbour-table OIDs through SNMP, and wait for a successful collector poll. "
    "A port with no LLDP/CDP advertisement is shown as link state only; it is not labelled as connected to an unknown device."
)

doc.add_heading("Collector ARP lookup", level=1)
doc.add_paragraph(
    "The selected device still controls authorization and profile selection, but the command itself reads the collector cache. "
    "This is intentional: an ARP table belongs to the collector host, and filtering it by 127.0.0.1 or another selected target hides the other learned entries."
)
doc.add_heading("Before", level=2)
code_paragraph(doc, "File: plugins/nms/includes/diagnostics.php\nLines 233-235\n\n$binary = nms_diag_program(\"ip\");\n$command = [$binary, \"neigh\", \"show\", \"to\", $target];")
doc.add_paragraph("This asked Linux for neighbours belonging only to the selected target. A successful command with no matching row produced an empty result.")
doc.add_heading("After", level=2)
code_paragraph(doc, "File: plugins/nms/includes/diagnostics.php\nLines 233-238\n\n$binary = nms_diag_program(\"ip\");\n$command = [$binary, \"neigh\", \"show\"];\n\nThe command now reads the complete collector neighbour cache.")
code_paragraph(doc, "File: plugins/nms/includes/diagnostics.php\nLines 287-291\n\nThe plugin runs only the configured ip neigh show command.\nThe result includes the executed command and reports when the cache is genuinely empty; no arp -an fallback is used.")
doc.add_paragraph("Readiness text is also updated in plugins/nms/diagnostics.php, lines 73-76. The run-page explanation is in plugins/nms/templates/diagnostics/run.php, line 33.")

doc.add_heading("iPerf3 bandwidth test", level=1)
doc.add_paragraph(
    "iPerf3 is a client and server pair. The collector runs the client, while the selected endpoint must run an iPerf3 server listening on TCP port 5201. "
    "A device being reachable by Ping or SNMP does not automatically mean that iPerf3 is available."
)
doc.add_heading("Before", level=2)
code_paragraph(doc, "File: plugins/nms/includes/diagnostics.php\nLines 247-249\n\n$binary = nms_diag_program(\"iperf3\");\n$command = [$binary, \"-c\", $target, \"-p\", \"5201\", \"-t\", $row[\"bandwidth_seconds\"], \"-J\"];\n\nThe client ran without checking whether a server was listening.")
doc.add_paragraph("With no server on 127.0.0.1:5201, the screenshot showed incomplete JSON and: iperf3: error - unable to send control message: Bad file descriptor.")
doc.add_heading("After", level=2)
code_paragraph(doc, "File: plugins/nms/includes/diagnostics.php\nLines 268-285\n\nThe plugin opens a short TCP check to $target:5201 before starting iperf3.\nIf it cannot connect, it returns exit code 111 and a plain explanation telling the operator to start an authorised iPerf3 server.\nIf the port is reachable, the original bounded iPerf3 command runs normally.")
doc.add_heading("Exact after code reference", level=2)
code_paragraph(doc, "File: plugins/nms/includes/diagnostics.php\nLines 268-285\n\n/* A bandwidth client needs a listening server before a measurement can begin. */\nif ($tool === \"iperf3\") {\n    $socket = @fsockopen($target, 5201, $socket_error, $socket_message, 2);\n    if (!is_resource($socket)) {\n        return [\"exit\" => 111, \"output\" => \"iPerf3 server is not reachable at \" . $target . \":5201...\"];\n    }\n    fclose($socket);\n}\n\n$result = nms_diag_run_command($command, 40);")
doc.add_paragraph("The readiness requirement remains visible in plugins/nms/diagnostics.php, lines 78-81: the remote endpoint must run an iPerf3 server on TCP 5201.")

doc.add_heading("How to use the new behavior", level=1)
doc.add_paragraph("For ARP, choose a device with an assigned diagnostic profile, select Collector ARP lookup, and run the test. The output represents the collector’s cache, not a remote device’s private ARP table.")
doc.add_paragraph("For iPerf3, start the server on the authorised remote endpoint:")
code_paragraph(doc, "iperf3 -s")
doc.add_paragraph("Then select the same device in Protocol checks and run iPerf3 bandwidth. If TCP 5201 is blocked or no server is running, the plugin now reports that condition directly.")

doc.add_heading("Netperf bandwidth test", level=1)
doc.add_paragraph(
    "Netperf uses a client on the Cacti collector and a netserver process on the selected endpoint. The plugin now checks TCP port 12865 before running the client, so a missing server produces a direct explanation instead of a confusing client error."
)
doc.add_heading("Installation without an RPM", level=2)
doc.add_paragraph(
    "If the offline RHEL host has no Netperf RPM, copy an approved Netperf source archive to the host, build it locally, and install it. The collector needs the netperf client; the target needs netserver."
)
code_paragraph(doc, "# On the offline RHEL collector, after copying netperf-<version>.tar.gz\ntar -xzf netperf-<version>.tar.gz\ncd netperf-<version>\n./configure --prefix=/usr/local\nmake -j$(nproc)\nsudo make install\nsudo /sbin/ldconfig\ncommand -v netperf\n\n# On the selected endpoint, start the server\ncommand -v netserver\nsudo netserver -D -p 12865")
doc.add_paragraph(
    "Open TCP 12865 between the collector and endpoint if a firewall blocks it. Then select Netperf bandwidth in Protocol checks and run the test. The plugin returns exit code 111 when the port is unavailable."
)
code_paragraph(doc, "File: plugins/nms/includes/diagnostics.php\nNetperf preflight added beside the iPerf3 check\n\n$port = $tool === \"iperf3\" ? 5201 : 12865;\n$socket = @fsockopen($target, $port, $socket_error, $socket_message, 2);\n\nNetperf requires netserver at $target:12865; iPerf3 requires an iPerf3 server at $target:5201.")

doc.add_heading("Validation", level=1)
doc.add_paragraph("The updated PHP files passed syntax checks on the Cacti RHEL VM. The diagnostics page returned HTTP 200 after deployment. The collector test confirmed that ip neigh show returns live IPv4 and IPv6 entries. The selected VM currently has no iPerf3 server listening on TCP 5201, so the new preflight correctly reports that requirement instead of running a failed client test.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("NMS plugin change reference")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(110, 110, 110)

doc.save(OUT)
print(OUT)
